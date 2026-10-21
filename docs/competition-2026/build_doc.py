"""Build the competition submission from the reviewed Markdown source."""
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parent
doc = Document()
for border in doc.styles.element.xpath('.//w:pBdr'):
    border.getparent().remove(border)
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(2)
sec.left_margin = sec.right_margin = Cm(2.2)
for name, size in [('Normal', 10.5), ('Title', 23), ('Heading 1', 15), ('Heading 2', 12)]:
    style = doc.styles[name]
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.2
doc.styles['Heading 1'].paragraph_format.space_before = Pt(10)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(6)
for line in (root / '作品说明.md').read_text(encoding='utf-8').splitlines():
    if not line.strip():
        continue
    if line == '---PAGE---':
        doc.add_page_break()
    elif line.startswith('# '):
        doc.add_paragraph(line[2:], 'Title')
    elif line.startswith('## '):
        doc.add_paragraph(line[3:], 'Heading 1')
    elif line.startswith('### '):
        doc.add_paragraph(line[4:], 'Heading 2')
    else:
        doc.add_paragraph(line)
footer = sec.footer.paragraphs[0]
footer.alignment = 2
footer.add_run('YOUAI 参赛作品说明  ·  ')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
for run in footer.runs:
    run.font.size = Pt(8)
doc.core_properties.title = 'YOUAI 微信小程序参赛作品说明'
doc.core_properties.subject = '2026微信小程序开发大赛作品说明'
doc.save(root / 'YOUAI微信小程序参赛作品说明.docx')
print('DOCX created')

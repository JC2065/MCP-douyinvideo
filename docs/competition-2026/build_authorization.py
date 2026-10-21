"""Build an unsigned, editable competition authorization form."""
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn

root = Path(__file__).resolve().parent
doc = Document()
for border in doc.styles.element.xpath('.//w:pBdr'):
    border.getparent().remove(border)
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(1.8)
sec.left_margin = sec.right_margin = Cm(2.2)
for name, size in [('Normal', 10.5), ('Title', 20)]:
    s = doc.styles[name]
    s.font.name = 'Microsoft YaHei'
    s.font.size = Pt(size)
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    s.paragraph_format.line_spacing = 1.15
    s.paragraph_format.space_after = Pt(8)
for line in (root / '参赛授权书.md').read_text(encoding='utf-8').splitlines():
    if not line.strip():
        continue
    if line.startswith('# '):
        p = doc.add_paragraph(line[2:], 'Title')
        p.alignment = 1
    else:
        doc.add_paragraph(line)
doc.core_properties.title = 'YOUAI参赛授权书'
doc.core_properties.subject = '通用拟稿 非官方模板 待填写签署'
doc.save(root / 'YOUAI参赛授权书.docx')
print('Authorization DOCX created')
